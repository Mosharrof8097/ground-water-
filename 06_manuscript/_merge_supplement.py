"""
Assemble the supplementary material.

    python _merge_supplement.py OUT.docx

The S-series tables are numbered here, independently of the main text, and the
table-drawing code is shared with the manuscript builder so both documents
format identically. Every table is read from 04_results/tables at build time.
"""

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

import importlib.util

HERE = Path(__file__).resolve().parent
SOURCE = HERE / "07_SUPPLEMENT.md"
TBL_DIR = HERE.parent / "04_results" / "tables"

_spec = importlib.util.spec_from_file_location("mm", HERE / "_merge_manuscript.py")
mm = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(mm)

# supplementary table key -> (file, caption without its number)
TABLES = {
    "S_CENSORING_CASES": ("S1_censoring_two_cases.csv",
                          "Two cases in which reading a detection limit as a "
                          "measurement creates structure that is not there."),
    "T2": ("T2_feasibility_matrix.csv",
           "Feasibility of every target and tier under random cross "
           "validation, with the model family selected inside the fold."),
    "T9": ("T9_spatial_vs_random_cv.csv",
           "Random against spatially blocked cross validation for every "
           "target, tier and block count."),
    "T0": ("T0_descriptive_statistics.csv",
           "Descriptive statistics for all predictors, targets and "
           "geochemistry, with censoring and detection limits."),
    "T3": ("T3_conformal_coverage.csv",
           "Held-out split conformal coverage and interval width for each "
           "feasible target and tier."),
    "T12": ("T12_feature_importance.csv",
            "Permutation feature importance for every modelled metal, "
            "measured as drop in log-scale R² on held-out folds."),
    "T13": ("T13_learning_curve.csv",
            "Learning curve results, held-out R² against training set size."),
    "T20": ("T20_log_skill_audit.csv",
            "Arsenic log-scale skill scored on all test wells and on measured "
            "test wells only, under identical folds."),
    "T4": ("T4_censoring_comparison.csv",
           "Four censoring treatments on identical folds against one common "
           "reference, with censored-consistency."),
    "T6": ("T6_charge_balance.csv",
           "Charge balance error statistics for the analytical quality check."),
    "T8": ("T8_water_types.csv",
           "Hydrochemical water types by Piper classification."),
    "T11": ("T11_depth_strata.csv",
            "Arsenic and iron by depth stratum."),
    "T14": ("T14_geology_arsenic.csv",
            "Arsenic by surface geology unit in the national survey."),
}


def main():
    out = Path(sys.argv[1])
    blocks = mm.read_markdown(SOURCE)

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = mm.FONT
    normal.font.size = Pt(mm.BODY_PT)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), mm.FONT)
    pf = normal.paragraph_format
    pf.line_spacing, pf.space_after = 1.5, Pt(0)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for section in doc.sections:
        mm.set_orientation(section, landscape=False)
    for level in (1, 2, 3):
        st = doc.styles[f"Heading {level}"]
        st.font.name, st.font.bold = mm.FONT, True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.font.size = Pt(mm.BODY_PT + (2 if level == 1 else 1 if level == 2 else 0))

    order = []
    for kind, level, payload in blocks:
        if kind == "heading":
            doc.add_heading(payload, level=min(level, 3))
        elif kind == "csv":
            key = payload
            if key not in TABLES:
                raise SystemExit(f"no supplementary table defined for {key}")
            name, caption = TABLES[key]
            order.append(key)
            path = TBL_DIR / name
            if not path.exists():
                raise SystemExit(f"missing {path}")
            import csv as _csv
            with path.open(newline="", encoding="utf-8") as fh:
                rows = [[mm.fmt(c) for c in r] for r in _csv.reader(fh)]
            rows[0] = [c.replace("_", " ") for c in rows[0]]
            n_cols = len(rows[0])
            wide = n_cols >= 13
            if wide:
                mm.set_orientation(doc.add_section(WD_SECTION.NEW_PAGE), True)
            mm.add_caption(doc, f"Table S{len(order)}. {caption}")
            mm.write_table(doc, rows,
                           size=7.0 if wide else 7.5 if n_cols >= 9 else mm.TABLE_PT,
                           avail=9.0 if wide else 6.5)
            if wide:
                mm.set_orientation(doc.add_section(WD_SECTION.NEW_PAGE), False)
        elif kind == "para":
            para = doc.add_paragraph()
            mm.add_runs(para, payload)

    doc.save(out)
    print(f"  saved {out.name}")
    print(f"  tables {len(doc.tables)} in order: "
          + ", ".join(f"S{i+1}={k}" for i, k in enumerate(order)))
    return 0


if __name__ == "__main__":
    sys.exit(main())
