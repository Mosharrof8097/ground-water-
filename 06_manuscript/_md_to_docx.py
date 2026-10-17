"""
Convert the assembled manuscript Markdown into a Word document formatted for
journal submission.

    python _md_to_docx.py INPUT.md OUTPUT.docx

Body text is Times New Roman 12 pt, double spaced, justified. Tables are set
single spaced in 9 pt with a bold header row so they stay legible at width, and
the reference list is given a hanging indent. The editorial notes section is
dropped, since it exists only for the authors.
"""

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

FONT = "Times New Roman"
BODY_PT = 12
TABLE_PT = 9


def set_cell_background(cell, hex_colour):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_colour)
    cell._tc.get_or_add_tcPr().append(shd)


def add_runs(paragraph, text):
    """Render **bold** and `code` spans; everything else is plain."""
    for part in re.split(r"(\*\*.+?\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Courier New"
            run.font.size = Pt(BODY_PT - 1)
        else:
            paragraph.add_run(part)


def parse_table(lines):
    """Markdown pipe table to a list of rows, discarding the separator line."""
    rows = []
    for line in lines:
        if re.match(r"^\s*\|[\s:|-]+\|\s*$", line):
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        rows.append(cells)
    return rows


def main():
    src, dst = Path(sys.argv[1]), Path(sys.argv[2])
    text = src.read_text(encoding="utf-8")

    # The editorial notes are for the authors, not the journal.
    text = text.split("## Editorial notes")[0]

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(BODY_PT)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    pf = normal.paragraph_format
    pf.line_spacing = 2.0
    pf.space_after = Pt(0)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for section in doc.sections:
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        for margin in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(section, margin, Inches(1))

    for level in (1, 2, 3):
        st = doc.styles[f"Heading {level}"]
        st.font.name = FONT
        st.font.size = Pt(BODY_PT + (2 if level == 1 else 1 if level == 2 else 0))
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)

    lines = text.split("\n")
    i = 0
    in_references = False

    while i < len(lines):
        line = lines[i].rstrip()

        if not line.strip() or line.strip() == "---":
            i += 1
            continue

        # Tables
        if line.lstrip().startswith("|"):
            block = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                block.append(lines[i])
                i += 1
            rows = parse_table(block)
            if not rows:
                continue
            n_cols = len(rows[0])
            table = doc.add_table(rows=len(rows), cols=n_cols)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False

            # Allocate width by the longest entry in each column, within
            # sensible bounds, so that narrow fields such as counts do not take
            # the same space as study names and headers stop breaking mid-word.
            usable = 6.5
            longest = [max(len(r[c]) for r in rows if c < len(r)) for c in range(n_cols)]
            weights = [min(max(w, 6), 26) for w in longest]
            total_w = sum(weights)
            widths = [Inches(usable * w / total_w) for w in weights]
            for c, width in enumerate(widths):
                for r in range(len(rows)):
                    table.cell(r, c).width = width

            for r, row in enumerate(rows):
                for c, cell_text in enumerate(row[: len(rows[0])]):
                    cell = table.cell(r, c)
                    cell.text = ""
                    para = cell.paragraphs[0]
                    para.paragraph_format.line_spacing = 1.0
                    para.paragraph_format.space_after = Pt(2)
                    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    add_runs(para, cell_text)
                    for run in para.runs:
                        run.font.size = Pt(TABLE_PT)
                        run.font.name = FONT
                        if r == 0:
                            run.bold = True
                    if r == 0:
                        set_cell_background(cell, "D9D9D9")
            doc.add_paragraph()
            continue

        # Headings
        m = re.match(r"^(#{1,3})\s+(.*)$", line)
        if m:
            level, title = len(m.group(1)), m.group(2).strip()
            in_references = title.lower().startswith("reference")
            doc.add_heading(title, level=level)
            i += 1
            continue

        # Gather the paragraph, which may wrap over several source lines
        block = []
        quoted = line.lstrip().startswith(">")
        while i < len(lines) and lines[i].strip() and not lines[i].lstrip().startswith("|") \
                and not re.match(r"^#{1,3}\s", lines[i]) and lines[i].strip() != "---":
            block.append(re.sub(r"^\s*>\s?", "", lines[i]).strip())
            i += 1
        para_text = " ".join(block).strip()
        if not para_text:
            continue

        para = doc.add_paragraph()
        if in_references:
            para.paragraph_format.line_spacing = 1.5
            para.paragraph_format.space_after = Pt(6)
            para.paragraph_format.left_indent = Inches(0.5)
            para.paragraph_format.first_line_indent = Inches(-0.5)
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif quoted:
            para.paragraph_format.left_indent = Inches(0.5)
            para.paragraph_format.right_indent = Inches(0.5)
            para.paragraph_format.space_before = Pt(6)
            para.paragraph_format.space_after = Pt(6)
            para.paragraph_format.line_spacing = 1.5
        elif para_text.startswith("**Table"):
            para.paragraph_format.line_spacing = 1.0
            para.paragraph_format.space_after = Pt(6)
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_runs(para, para_text)
        if quoted:
            for run in para.runs:
                run.italic = True

    doc.save(dst)
    print(f"saved {dst.name}")
    print(f"  paragraphs {len(doc.paragraphs)}, tables {len(doc.tables)}")


if __name__ == "__main__":
    main()
